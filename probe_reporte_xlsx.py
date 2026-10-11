import requests

s = requests.Session()
s.post("http://127.0.0.1:5001/login",
       data={"correo": "admin@cleoferr.com", "clave": "Admin123!"},
       allow_redirects=False)

# 1) Generar: debe renderizar la tabla Resultados (HTML), NO descargar archivo
r = s.get("http://127.0.0.1:5002/reportes/generar?tipo=ventas&desde=2026-01-01&hasta=2026-12-31")
html = "text/html" in r.headers.get("Content-Type", "")
tiene_resultados = "Resultados" in r.text
tiene_export_xlsx = ".xlsx" in r.text and "export_url" not in r.headers.get("Content-Disposition", "")
print(f"GET /reportes/generar -> {r.status_code} html={html} tabla_resultados={tiene_resultados} boton_xlsx={'/reportes/ventas/exportar' in r.text}")

# 2) Exportar: debe descargar un .xlsx real
r2 = s.get("http://127.0.0.1:5002/reportes/ventas/exportar?desde=2026-01-01&hasta=2026-12-31")
es_xlsx = r2.headers.get("Content-Type", "").startswith("application/vnd.openxmlformats")
cd = r2.headers.get("Content-Disposition", "")
print(f"GET /reportes/ventas/exportar -> {r2.status_code} xlsx={es_xlsx} {cd} bytes={len(r2.content)}")

# Validacion rapida del contenido xlsx (magic bytes PK + legible por openpyxl)
if es_xlsx:
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r2.content))
    ws = wb.active
    print(f"xlsx hoja='{ws.title}' filas={ws.max_row} cols={ws.max_column} header={[c.value for c in ws[1]]}")
